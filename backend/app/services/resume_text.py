"""Turn an uploaded resume file (PDF, DOCX or TXT) into plain text."""

import io
import logging
import re

logger = logging.getLogger(__name__)

MAX_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 10
MIN_TEXT_CHARS = 50
ALLOWED_EXTENSIONS = (".pdf", ".docx", ".txt")


class ResumeFileError(ValueError):
    """The file cannot be used as a resume; the message is safe to show the user."""


def _tidy(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _pdf_text(data: bytes) -> str:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n\n".join((page.extract_text() or "") for page in pdf.pages[:MAX_PDF_PAGES])


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def extract_resume_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if not name.endswith(ALLOWED_EXTENSIONS):
        raise ResumeFileError("Unsupported file type. Upload a PDF, DOCX or TXT resume.")
    if len(data) > MAX_BYTES:
        raise ResumeFileError(f"File is larger than {MAX_BYTES // (1024 * 1024)} MB.")
    if not data:
        raise ResumeFileError("The file is empty.")

    # Check the real content, not just the extension.
    if name.endswith(".pdf") and not data.startswith(b"%PDF"):
        raise ResumeFileError("This file is not a valid PDF.")
    if name.endswith(".docx") and not data.startswith(b"PK"):
        raise ResumeFileError("This file is not a valid DOCX.")

    try:
        if name.endswith(".pdf"):
            raw = _pdf_text(data)
        elif name.endswith(".docx"):
            raw = _docx_text(data)
        else:
            raw = data.decode("utf-8", errors="replace")
    except Exception as exc:  # pdfplumber and python-docx raise many unrelated exception types
        logger.warning("could not read %s: %s", filename, type(exc).__name__)
        raise ResumeFileError("Could not read this file. It may be damaged or password-protected.") from exc

    text = _tidy(raw)
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeFileError(
            "No readable text found. Scanned or image-only resumes are not supported (no OCR)."
        )
    return text
