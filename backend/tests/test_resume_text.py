import io

import pytest
from docx import Document
from fpdf import FPDF

from app.services.resume_text import MAX_BYTES, ResumeFileError, extract_resume_text

BODY = "Asha Verma. Backend engineer with five years of Python, FastAPI and PostgreSQL experience."


def make_pdf(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    if text:
        pdf.multi_cell(0, 8, text)
    return bytes(pdf.output())


def make_docx(paragraph: str, table_cell: str | None = None) -> bytes:
    doc = Document()
    doc.add_paragraph(paragraph)
    if table_cell:
        doc.add_table(rows=1, cols=2).rows[0].cells[0].text = table_cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extracts_text_from_pdf():
    assert "FastAPI" in extract_resume_text("cv.pdf", make_pdf(BODY))


def test_extracts_text_from_docx_including_tables():
    text = extract_resume_text("cv.docx", make_docx(BODY, "Skills: Python, SQL"))
    assert "PostgreSQL" in text and "Skills: Python, SQL" in text


def test_extracts_text_from_txt():
    assert extract_resume_text("cv.txt", BODY.encode()) == BODY


def test_extension_check_is_case_insensitive():
    assert "Asha" in extract_resume_text("CV.PDF", make_pdf(BODY))


@pytest.mark.parametrize("name", ["cv.exe", "cv", "cv.png", "cv.pdf.exe"])
def test_rejects_unsupported_types(name):
    with pytest.raises(ResumeFileError, match="Unsupported"):
        extract_resume_text(name, b"x" * 100)


def test_rejects_files_whose_content_does_not_match_the_extension():
    with pytest.raises(ResumeFileError, match="not a valid PDF"):
        extract_resume_text("cv.pdf", b"<html>not a pdf</html>" * 10)
    with pytest.raises(ResumeFileError, match="not a valid DOCX"):
        extract_resume_text("cv.docx", b"plain text pretending to be docx" * 5)


def test_rejects_corrupt_files_with_a_friendly_message():
    with pytest.raises(ResumeFileError, match="Could not read"):
        extract_resume_text("cv.pdf", b"%PDF-1.4 this is not really a pdf")


def test_scanned_or_blank_pdf_gives_a_clear_no_text_error():
    with pytest.raises(ResumeFileError, match="No readable text"):
        extract_resume_text("cv.pdf", make_pdf(""))


def test_rejects_empty_and_oversized_files():
    with pytest.raises(ResumeFileError, match="empty"):
        extract_resume_text("cv.txt", b"")
    with pytest.raises(ResumeFileError, match="larger"):
        extract_resume_text("cv.txt", b"a" * (MAX_BYTES + 1))


def test_whitespace_is_tidied():
    text = extract_resume_text("cv.txt", (BODY + "\n\n\n\n\n" + "second   line   here " * 3).encode())
    assert "\n\n\n" not in text and "  " not in text
